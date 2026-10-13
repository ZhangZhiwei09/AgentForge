"""End-to-end file/text ingestion, atomic persistence and worker contracts."""

import asyncio
import io
import os
import subprocess
import sys
import uuid
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
import pytest_asyncio
from arq import Retry, create_pool
from arq.connections import RedisSettings
from docx import Document
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import func, select, update

from src.config import settings
from src.jobs import ingestion as jobs
from src.models.base import Base
from src.models.knowledge import KnowledgeBase, KnowledgeChunk, KnowledgeDocument
from src.schemas.chat import ChatSyncResult
from src.services import (
    document_parser,
    knowledge_ingestion,
    knowledge_storage,
)
from src.services.knowledge_ingestion import IngestionService
from src.worker import IngestionWorker

USER = {"X-Test-Actor": "user"}
CONTENT = "核身接口超时需要检查网络连接、权限设置、请求参数和服务器日志。\n\n" * 20


def test_worker_imports_in_fresh_interpreter():
    result = subprocess.run(
        [sys.executable, "-c", "from src.worker import WorkerSettings; assert WorkerSettings.functions"],
        capture_output=True, text=True, timeout=30, check=False,
    )
    assert result.returncode == 0, result.stderr


@pytest_asyncio.fixture
async def ingestion_api(entry_api, monkeypatch, tmp_path):
    client, sessions = entry_api
    async with sessions.kw["bind"].begin() as connection:
        await connection.run_sync(lambda sync: Base.metadata.create_all(sync, tables=[
            KnowledgeBase.__table__, KnowledgeDocument.__table__, KnowledgeChunk.__table__,
        ]))
    embed = AsyncMock(side_effect=lambda texts: [[0.1] * 1024 for _ in texts])
    monkeypatch.setattr(knowledge_ingestion, "get_embedding_provider", lambda: SimpleNamespace(embed=embed))
    monkeypatch.setattr(settings, "ingestion_async_enabled", False)
    monkeypatch.setattr(settings, "elasticsearch_url", "")
    monkeypatch.setattr(settings, "minio_endpoint", "")
    monkeypatch.setattr(settings, "knowledge_storage_root", str(tmp_path / "storage"))
    response = await client.post("/api/knowledge/bases", headers=USER, json={
        "name": "摄入知识", "chunk_size_tokens": 200, "chunk_overlap_tokens": 0,
    })
    assert response.status_code == 201, response.text
    return client, sessions, response.json()["id"], embed


@pytest.mark.asyncio
async def test_text_creation_and_legacy_ingest_share_pipeline(ingestion_api):
    client, sessions, base_id, embed = ingestion_api
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents", headers=USER, json={
        "title": "超时排查", "content": CONTENT,
    })
    assert response.status_code == 201, response.text
    result = response.json()
    assert result["status"] == "completed"
    assert result["chunkCount"] > 1
    assert result["processingDetail"]["phase"] == "completed"
    for field in ("processingStartedAt", "parsingCompletedAt", "normalizingCompletedAt",
                  "chunkingCompletedAt", "embeddingCompletedAt"):
        assert result[field]
    legacy = await client.post("/api/knowledge/ingest", headers=USER, json={
        "kb_id": base_id, "title": "兼容摄入", "content": CONTENT, "chunk_size": 100, "chunk_overlap": 0,
    })
    assert legacy.status_code == 200, legacy.text
    assert legacy.json()["status"] == "completed"
    async with sessions() as db:
        chunks = list(await db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == legacy.json()["doc_id"])))
        assert all(chunk.token_count <= 100 for chunk in chunks)
        assert all(len(chunk.embedding) == 1024 for chunk in chunks)
    assert embed.await_count == 2


@pytest.mark.asyncio
async def test_upload_process_false_waits_for_saved_settings(ingestion_api):
    client, sessions, base_id, embed = ingestion_api
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents/upload?process=false", headers=USER,
                                 files={"file": ("guide.md", CONTENT.encode(), "text/markdown")})
    assert response.status_code == 201, response.text
    document = response.json()
    assert document["status"] == "pending"
    assert document["file_name"] == "guide.md"
    assert document["file_size"] == len(CONTENT.encode())
    embed.assert_not_awaited()
    assert (await client.put(f"/api/knowledge/bases/{base_id}", headers=USER, json={
        "chunk_size_tokens": 100, "chunk_overlap_tokens": 0,
    })).status_code == 200
    processed = await client.post(f"/api/knowledge/bases/{base_id}/process", headers=USER)
    assert processed.json()["processed"] == 1
    async with sessions() as db:
        saved = await db.get(KnowledgeDocument, document["id"])
        assert saved.status == "completed"
        chunks = list(await db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document["id"])))
        assert all(chunk.token_count <= 100 for chunk in chunks)
    assert (await client.post(f"/api/knowledge/bases/{base_id}/process", headers=USER)).json()["processed"] == 0


@pytest.mark.asyncio
async def test_docx_upload_is_binary_and_parsed_in_pipeline(ingestion_api):
    client, _, base_id, _ = ingestion_api
    document = Document()
    document.add_heading("接口超时排查", 1)
    document.add_paragraph(CONTENT)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "错误码"
    table.cell(0, 1).text = "方案"
    table.cell(1, 0).text = "TIMEOUT"
    table.cell(1, 1).text = "检查网络"
    output = io.BytesIO()
    document.save(output)
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents/upload", headers=USER, files={
        "file": ("guide.docx", output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    })
    assert response.status_code == 201, response.text
    result = response.json()
    assert result["status"] == "completed"
    assert result["content"].startswith("# 接口超时排查")
    assert "| 错误码 | 方案 |" in result["content"]
    assert result["downloadingCompletedAt"]
    assert result["originalFilePath"]
    assert await knowledge_storage.read_file(result["originalFilePath"]) == output.getvalue()


@pytest.mark.asyncio
async def test_hierarchical_ingestion_vectors_only_children(ingestion_api):
    client, sessions, base_id, embed = ingestion_api
    assert (await client.put(f"/api/knowledge/bases/{base_id}", headers=USER, json={
        "chunk_structure": "hierarchical", "child_chunk_size_tokens": 50, "child_chunk_overlap_tokens": 0,
    })).status_code == 200
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents", headers=USER, json={
        "title": "父子分块", "content": CONTENT,
    })
    assert response.status_code == 201
    assert response.json()["status"] == "completed"
    async with sessions() as db:
        chunks = list(await db.scalars(select(KnowledgeChunk)))
        parents = [chunk for chunk in chunks if chunk.parent_chunk_id is None]
        children = [chunk for chunk in chunks if chunk.parent_chunk_id is not None]
        assert len(parents) > 1
        assert len(children) > len(parents)
        assert len({chunk.chunk_index for chunk in chunks}) == len(chunks)
        assert all(chunk.embedding is None for chunk in parents)
        assert all(len(chunk.embedding) == 1024 for chunk in children)
        assert response.json()["chunkCount"] == len(chunks)
        assert sum(len(call.args[0]) for call in embed.await_args_list) == len(children)


@pytest.mark.asyncio
async def test_embedding_failure_preserves_prior_chunks_and_can_retry(ingestion_api):
    _, sessions, base_id, embed = ingestion_api
    async with sessions() as db:
        db.add(KnowledgeDocument(
            id="retry", knowledge_base_id=base_id, title="retry", content=CONTENT, status="pending", chunk_count=1,
        ))
        await db.flush()
        db.add(KnowledgeChunk(id="old", document_id="retry", knowledge_base_id=base_id,
                              chunk_index=0, content="旧分块", token_count=4))
        await db.commit()
    embed.side_effect = RuntimeError("embedding unavailable")
    async with sessions() as db:
        result = await IngestionService(db).process("retry", base_id)
        assert result["status"] == "failed"
    async with sessions() as db:
        assert await db.get(KnowledgeChunk, "old") is not None
        assert (await db.get(KnowledgeDocument, "retry")).status == "failed"
    embed.side_effect = lambda texts: [[0.2] * 1024 for _ in texts]
    async with sessions() as db:
        result = await IngestionService(db).process("retry", base_id)
        assert result["status"] == "completed"
    async with sessions() as db:
        assert await db.get(KnowledgeChunk, "old") is None
        assert (await db.get(KnowledgeDocument, "retry")).retry_count == 1


@pytest.mark.asyncio
async def test_batch_failure_does_not_expire_following_documents(ingestion_api):
    client, sessions, base_id, embed = ingestion_api
    embed.side_effect = RuntimeError("offline")
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents/batch", headers=USER, json={
        "documents": [{"title": "a", "content": CONTENT}, {"title": "b", "content": CONTENT}],
    })
    assert response.status_code == 200, response.text
    assert len(response.json()["doc_ids"]) == 2
    async with sessions() as db:
        assert list(await db.scalars(select(KnowledgeDocument.status))) == ["failed", "failed"]


@pytest.mark.asyncio
async def test_overlapping_worker_claims_process_once(ingestion_api):
    _, sessions, base_id, embed = ingestion_api
    async with sessions() as db:
        db.add(KnowledgeDocument(id="concurrent", knowledge_base_id=base_id, title="task", content=CONTENT))
        await db.commit()
    reached = asyncio.Event()
    release = asyncio.Event()

    async def waiting_embed(texts):
        reached.set()
        await release.wait()
        return [[0.1] * 1024 for _ in texts]

    embed.side_effect = waiting_embed

    async def process():
        async with sessions() as db:
            return await IngestionService(db).process("concurrent", base_id)

    running = asyncio.create_task(process())
    await asyncio.wait_for(reached.wait(), 5)
    second = await process()
    assert second["status"] == "busy"
    release.set()
    assert (await running)["status"] == "completed"
    assert (await process())["processed"] is False
    embed.assert_awaited_once()


@pytest.mark.asyncio
async def test_expired_claim_is_recovered(ingestion_api):
    _, sessions, base_id, _ = ingestion_api
    async with sessions() as db:
        db.add(KnowledgeDocument(
            id="stale", knowledge_base_id=base_id, title="stale", content=CONTENT, status="embedding",
            processing_started_at=datetime.now(UTC) - timedelta(hours=1), updated_at=datetime.now(UTC) - timedelta(hours=1),
        ))
        await db.commit()
        result = await IngestionService(db).process("stale", base_id)
    assert result["status"] == "completed"


@pytest.mark.asyncio
async def test_superseded_worker_cannot_replace_new_chunks(ingestion_api):
    _, sessions, base_id, embed = ingestion_api
    async with sessions() as db:
        db.add(KnowledgeDocument(id="takeover", knowledge_base_id=base_id, title="task", content=CONTENT))
        await db.commit()
    reached = asyncio.Event()
    release = asyncio.Event()

    async def waiting_embed(texts):
        reached.set()
        await release.wait()
        return [[0.1] * 1024 for _ in texts]

    embed.side_effect = waiting_embed

    async def process():
        async with sessions() as db:
            return await IngestionService(db).process("takeover", base_id)

    running = asyncio.create_task(process())
    try:
        await asyncio.wait_for(reached.wait(), 5)
        async with sessions() as db:
            await db.execute(update(KnowledgeDocument).where(KnowledgeDocument.id == "takeover").values(
                updated_at=datetime.now(UTC) - timedelta(hours=1),
            ))
            await db.commit()
        embed.side_effect = lambda texts: [[0.2] * 1024 for _ in texts]
        assert (await process())["status"] == "completed"
    finally:
        release.set()
    assert (await running)["status"] == "busy"
    async with sessions() as db:
        assert (await db.get(KnowledgeDocument, "takeover")).status == "completed"
        chunks = list(await db.scalars(select(KnowledgeChunk)))
        assert chunks
        assert all(abs(chunk.embedding[0] - 0.2) < 0.00001 for chunk in chunks)


@pytest.mark.asyncio
async def test_scan_pdf_and_invalid_doc_fail_without_chunks(ingestion_api):
    client, sessions, base_id, embed = ingestion_api
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    for filename, content, mime in [
        ("blank.pdf", output.getvalue(), "application/pdf"),
        ("legacy.doc", b"binary legacy doc", "application/msword"),
    ]:
        response = await client.post(f"/api/knowledge/bases/{base_id}/documents/upload", headers=USER,
                                     files={"file": (filename, content, mime)})
        assert response.status_code == 201
        assert response.json()["status"] == "failed"
        assert response.json()["errorMessage"]
    embed.assert_not_awaited()
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(KnowledgeChunk)) == 0


@pytest.mark.asyncio
async def test_image_ocr_reuses_existing_provider(ingestion_api, monkeypatch):
    client, _, base_id, _ = ingestion_api
    chat = AsyncMock(return_value=ChatSyncResult(content=CONTENT, usage={}))
    monkeypatch.setattr(document_parser, "resolve_model", lambda model: {"provider_name": "test", "model_id": model})
    monkeypatch.setattr(document_parser, "get_provider", lambda name: SimpleNamespace(chat_sync=chat))
    image = Image.new("RGB", (8, 8), "white")
    output = io.BytesIO()
    image.save(output, format="PNG")
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents/upload", headers=USER,
                                 files={"file": ("ocr.png", output.getvalue(), "image/png")})
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "completed"
    message = chat.await_args.args[0][0]
    assert message.content[0]["image_url"]["url"].startswith("data:image/png;base64,")


@pytest.mark.asyncio
async def test_upload_validation_and_path_traversal(ingestion_api):
    client, _, base_id, embed = ingestion_api
    url = f"/api/knowledge/bases/{base_id}/documents/upload"
    assert (await client.post(url, headers=USER)).status_code == 400
    assert (await client.post(url, headers=USER, files={"file": ("x.exe", b"binary", "application/octet-stream")})).status_code == 400
    assert (await client.post(url, headers=USER, files={"file": ("x.txt", b"   ", "text/plain")})).status_code == 400
    assert (await client.post(url, files={"file": ("x.txt", CONTENT.encode(), "text/plain")})).status_code == 401
    response = await client.post(url + "?process=false", headers=USER,
                                 files={"file": ("../../guide.txt", CONTENT.encode(), "text/plain")})
    assert response.status_code == 201
    assert response.json()["originalFilename"] == "guide.txt"
    with pytest.raises(ValueError):
        await knowledge_storage.save_file("../outside", b"data")
    with pytest.raises(ValueError):
        knowledge_storage.object_key("a/../outside")
    embed.assert_not_awaited()


@pytest.mark.asyncio
async def test_queue_dispatch_is_deduplicated_and_does_not_process_in_api(ingestion_api, monkeypatch):
    client, sessions, base_id, embed = ingestion_api
    pool = SimpleNamespace(enqueue_job=AsyncMock(), aclose=AsyncMock())
    create = AsyncMock(return_value=pool)
    monkeypatch.setattr(jobs, "create_pool", create)
    monkeypatch.setattr(settings, "ingestion_async_enabled", True)
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents", headers=USER, json={
        "title": "queued", "content": CONTENT,
    })
    assert response.status_code == 201, response.text
    document_id = response.json()["id"]
    assert response.json()["status"] == "pending"
    assert pool.enqueue_job.await_args.kwargs["_job_id"] == f"ingest-{document_id}"
    pool.aclose.assert_awaited_once()
    embed.assert_not_awaited()
    result = await jobs.ingest_document_job({"session_factory": sessions, "job_try": 1}, document_id, base_id)
    assert result["status"] == "completed"
    embed.assert_awaited_once()


@pytest.mark.asyncio
async def test_queue_outage_preserves_pending_document(ingestion_api, monkeypatch):
    client, sessions, base_id, embed = ingestion_api
    monkeypatch.setattr(settings, "ingestion_async_enabled", True)
    monkeypatch.setattr(jobs, "create_pool", AsyncMock(side_effect=ConnectionError("offline")))
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents", headers=USER, json={
        "title": "queued", "content": CONTENT,
    })
    assert response.status_code == 503
    async with sessions() as db:
        assert list(await db.scalars(select(KnowledgeDocument.status))) == ["pending"]
    embed.assert_not_awaited()


@pytest.mark.asyncio
async def test_worker_retry_and_deleted_document_contracts(ingestion_api):
    _, sessions, base_id, embed = ingestion_api
    async with sessions() as db:
        db.add(KnowledgeDocument(id="retry-job", knowledge_base_id=base_id, title="retry", content=CONTENT))
        await db.commit()
    embed.side_effect = RuntimeError("offline")
    with pytest.raises(Retry):
        await jobs.ingest_document_job({"session_factory": sessions, "job_try": 1}, "retry-job", base_id)
    assert (await jobs.ingest_document_job({"session_factory": sessions, "job_try": 3}, "retry-job", base_id))["status"] == "failed"
    assert (await jobs.ingest_document_job({"session_factory": sessions}, "deleted", base_id))["status"] == "deleted"


@pytest.mark.asyncio
async def test_arq_worker_consumes_real_redis_job(ingestion_api):
    redis_url = os.environ.get("AGENTFORGE_TEST_REDIS_URL")
    if not redis_url:
        pytest.skip("Set AGENTFORGE_TEST_REDIS_URL for isolated Redis worker integration")
    _, sessions, base_id, embed = ingestion_api
    document_id = str(uuid.uuid4())
    async with sessions() as db:
        db.add(KnowledgeDocument(id=document_id, knowledge_base_id=base_id, title="queue", content=CONTENT))
        await db.commit()
    queue = f"agentforge:test:{uuid.uuid4().hex}"
    pool = await create_pool(RedisSettings.from_dsn(redis_url), default_queue_name=queue)
    await pool.enqueue_job("ingest_document_job", document_id, base_id, _job_id=document_id)
    worker = IngestionWorker(
        [jobs.ingest_document_job], redis_pool=pool, queue_name=queue,
        burst=True, handle_signals=False, keep_result=0, ctx={"session_factory": sessions},
    )
    try:
        await worker.async_run()
        assert worker.jobs_complete == 1
        embed.assert_awaited_once()
        async with sessions() as db:
            assert (await db.get(KnowledgeDocument, document_id)).status == "completed"
    finally:
        await worker.close()


@pytest.mark.asyncio
async def test_pdf_text_extraction_and_encryption_errors(ingestion_api):
    client, _, base_id, _ = ingestion_api
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=200)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    stream = DecodedStreamObject()
    text = "Network timeout requires checking permissions and server logs. " * 5
    stream.set_data(f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents/upload", headers=USER,
                                 files={"file": ("text.pdf", output.getvalue(), "application/pdf")})
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "completed"
    assert "Network timeout" in response.json()["content"]
    writer.encrypt("password")
    output = io.BytesIO()
    writer.write(output)
    response = await client.post(f"/api/knowledge/bases/{base_id}/documents/upload", headers=USER,
                                 files={"file": ("encrypted.pdf", output.getvalue(), "application/pdf")})
    assert response.status_code == 201
    assert response.json()["status"] == "failed"
    assert "已加密" in response.json()["errorMessage"]


@pytest.mark.asyncio
async def test_cancelled_ingestion_records_failure_for_worker_retry(ingestion_api):
    _, sessions, base_id, embed = ingestion_api
    async with sessions() as db:
        db.add(KnowledgeDocument(id="cancel", knowledge_base_id=base_id, title="cancel", content=CONTENT))
        await db.commit()
    reached = asyncio.Event()

    async def waiting_embed(texts):
        reached.set()
        await asyncio.Future()

    embed.side_effect = waiting_embed

    async def process():
        async with sessions() as db:
            return await IngestionService(db).process("cancel", base_id)

    running = asyncio.create_task(process())
    await asyncio.wait_for(reached.wait(), 5)
    running.cancel()
    with pytest.raises(asyncio.CancelledError):
        await running
    async with sessions() as db:
        document = await db.get(KnowledgeDocument, "cancel")
        assert document.status == "failed"
        assert "处理已中断" in document.error_message
        assert await db.scalar(select(func.count()).select_from(KnowledgeChunk)) == 0


@pytest.mark.asyncio
async def test_legacy_ingest_preserves_404_and_rejects_invalid_overrides(ingestion_api):
    client, sessions, base_id, embed = ingestion_api
    response = await client.post("/api/knowledge/ingest", headers=USER, json={
        "kb_id": "absent", "title": "missing", "content": CONTENT,
    })
    assert response.status_code == 404
    response = await client.post("/api/knowledge/ingest", headers=USER, json={
        "kb_id": base_id, "title": "invalid", "content": CONTENT, "chunk_size": 100, "chunk_overlap": 60,
    })
    assert response.status_code == 422
    embed.assert_not_awaited()
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(KnowledgeDocument)) == 0
