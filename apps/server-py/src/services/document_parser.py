"""File-type selection and extraction using maintained format libraries."""

import asyncio
import base64
import io
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image
from pypdf import PdfReader

from src.config import settings
from src.providers.registry import get_provider, resolve_model
from src.schemas.chat import ChatMessage

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {"txt", "md", "json", "csv", "html", "xml", "yaml", "yml", "log"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "tiff", "tif"}


@dataclass
class ParsedDocument:
    text: str
    parser_name: str
    page_count: int | None = None


def parser_name(filename: str, mime_type: str | None = None) -> str | None:
    extension = Path(filename).suffix.lstrip(".").lower()
    mime = (mime_type or "").lower().split(";")[0]
    if extension == "pdf" or mime == "application/pdf":
        return "pdf"
    if extension in {"doc", "docx"} or mime in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword",
    }:
        return "word"
    if extension in IMAGE_EXTENSIONS or mime.startswith("image/"):
        return "image"
    if extension in TEXT_EXTENSIONS or mime.startswith((
        "text/", "application/json", "application/xml", "application/x-yaml",
    )):
        return "text"
    return None


def _parse_pdf(content: bytes) -> ParsedDocument:
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ValueError("PDF 已加密，请上传未设置密码保护的版本")
    if len(reader.pages) > 2000:
        raise ValueError("PDF 页数超过处理上限")
    return ParsedDocument("\n\n".join(page.extract_text() or "" for page in reader.pages), "pdf", len(reader.pages))


def _parse_word(content: bytes) -> ParsedDocument:
    if not content.startswith(b"PK"):
        raise ValueError("旧版 .doc 请转换为 .docx 后上传")
    with ZipFile(io.BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Word 解压后超过处理上限")
    document = Document(io.BytesIO(content))
    parts = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else ""
            heading = re.fullmatch(r"Heading ([1-6])", style)
            prefix = "#" * int(heading[1]) + " " if heading else "- " if style.startswith("List") else ""
            if block.text.strip():
                parts.append(prefix + block.text)
        elif isinstance(block, Table):
            rows = ["| " + " | ".join(cell.text.replace("|", "\\|").replace("\n", "<br>") for cell in row.cells) + " |"
                    for row in block.rows]
            if rows:
                rows.insert(1, "| " + " | ".join("---" for _ in block.rows[0].cells) + " |")
                parts.append("\n".join(rows))
    return ParsedDocument("\n\n".join(parts), "word")


async def _parse_image(content: bytes) -> ParsedDocument:
    def image_mime():
        with Image.open(io.BytesIO(content)) as image:
            image.verify()
            return Image.MIME.get(image.format, "image/png")

    mime = await asyncio.to_thread(image_mime)
    try:
        resolved = resolve_model(settings.video_model)
        provider = get_provider(resolved["provider_name"])
        result = await provider.chat_sync(
            [ChatMessage(role="user", content=[
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{base64.b64encode(content).decode('ascii')}"}},
                {"type": "text", "text": "请完整提取图片中的文字，保留表格结构，并简要描述图表或照片内容。只输出 Markdown 内容，不添加额外说明。"},
            ])],
            model=resolved["model_id"], temperature=0.1, max_tokens=2000,
        )
        if result.content.strip():
            return ParsedDocument(result.content, "image")
    except Exception:
        logger.exception("Image OCR unavailable")
    return ParsedDocument(f"[图片文件: {mime}, 大小: {round(len(content) / 1024)}KB, 未提取文字内容]", "image")


async def parse_document(content: bytes, filename: str, mime_type: str | None = None) -> ParsedDocument:
    name = parser_name(filename, mime_type)
    if name == "pdf":
        return await asyncio.to_thread(_parse_pdf, content)
    if name == "word":
        return await asyncio.to_thread(_parse_word, content)
    if name == "image":
        return await _parse_image(content)
    if name == "text":
        return ParsedDocument(content.decode("utf-8-sig", errors="replace"), "text")
    raise ValueError("不支持的文件类型")


def normalize_text(content: str, *, remove_extra_spaces: bool, remove_urls_emails: bool) -> str:
    text = content.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if remove_urls_emails:
        text = re.sub(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", "", text)
        text = re.sub(r"(?:https?|ftp)://[^\s)]+", "", text)
    if remove_extra_spaces:
        text = "\n".join(re.sub(r"[ \t]{2,}", " ", line).strip() for line in text.split("\n"))
    return text.strip()
